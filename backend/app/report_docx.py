"""The counter-proposal report as an editable DOCX (A4 portrait), from report.report_blocks.

The same content as the PDF, as text and tables the freelancer can edit before sending.
Generated in memory, never stored.
"""

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from app.report import REPORT_FOOTER, Block

A4 = (Mm(210), Mm(297))
PROTOCOL_WIDTHS = (Mm(57), Mm(56), Mm(56))
SUMMARY_WIDTHS = (Mm(40), Mm(25))


def render_docx(blocks: list[Block]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = A4
    section.left_margin = section.right_margin = Mm(20)
    document.styles["Normal"].font.size = Pt(10)
    document.core_properties.title = blocks[0].text if blocks else "Counter-proposal"
    document.core_properties.author = "Anti-Scope Creep"
    _footer(section.footer.paragraphs[0])
    for block in blocks:
        _render(document, block)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render(document, block: Block) -> None:
    if block.kind == "title":
        document.add_heading(block.text, level=0)
    elif block.kind == "heading":
        document.add_heading(block.text, level=1)
    elif block.kind == "paragraph":
        _add_lines(document.add_paragraph(), block.text)
    else:
        _table(document, block)


def _table(document, block: Block) -> None:
    widths = PROTOCOL_WIDTHS if len(block.header) == 3 else SUMMARY_WIDTHS
    table = document.add_table(rows=1, cols=len(block.header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, block.header):
        _write(cell, text, bold=True)
    for values in block.rows:
        for cell, text in zip(table.add_row().cells, values):
            _write(cell, text)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width
    document.add_paragraph()


def _write(cell: _Cell, text: str, bold: bool = False) -> None:
    _add_lines(cell.paragraphs[0], text, bold=bold)


def _add_lines(paragraph: Paragraph, text: str, bold: bool = False) -> None:
    """Keep line breaks (and so the email's bullets) as breaks inside one paragraph."""
    for index, line in enumerate(text.split("\n")):
        run = paragraph.add_run(line)
        run.bold = bold
        if index < text.count("\n"):
            run.add_break()


def _footer(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, line in enumerate(REPORT_FOOTER):
        run = paragraph.add_run(line)
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(0x6E, 0x6E, 0x6E)
        if index < len(REPORT_FOOTER) - 1:
            run.add_break()
